"""Harmonize South Korean municipal mortality data for ABI feasibility checks."""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import geopandas as gpd
import networkx as nx
import numpy as np
import openpyxl
import pandas as pd
from shapely import make_valid


BASE_DIR = Path(__file__).resolve().parent
RAW_DIR = BASE_DIR / "raw"
YEARS = range(2015, 2020)

CAUSES = {
    "lung_cancer": "Malignant neoplasms of trachea, bronchus and lung (C33-C34)",
    "suicide": "Intentional self-harm (X60-X84)",
    "circulatory": "Diseases of the circulatory system (I00-I99)",
}

OLD_TO_CURRENT_PROVINCE = {
    "11": "11",
    "21": "26",
    "22": "27",
    "23": "28",
    "24": "29",
    "25": "30",
    "26": "31",
    "29": "36",
    "31": "41",
    "32": "42",
    "33": "43",
    "34": "44",
    "35": "45",
    "36": "46",
    "37": "47",
    "38": "48",
    "39": "50",
}

# KOSIS population ages are collapsed to the detailed mortality age groups.
POPULATION_AGE_MAP = {
    "020": "00-04",
    "050": "05-09",
    "070": "10-14",
    "100": "15-19",
    "120": "20-24",
    "130": "25-29",
    "150": "30-34",
    "160": "35-39",
    "180": "40-44",
    "190": "45-49",
    "210": "50-54",
    "230": "55-59",
    "260": "60-64",
    "280": "65-69",
    "310": "70-74",
    "330": "75-79",
    "360": "80-84",
    "380": "85-89",
    "410": "90+",
    "430": "90+",
    "440": "90+",
}

MORTALITY_AGE_MAP = {
    "03": "00-04",
    "04": "00-04",
    "05": "05-09",
    "10": "10-14",
    "15": "15-19",
    "20": "20-24",
    "25": "25-29",
    "30": "30-34",
    "35": "35-39",
    "40": "40-44",
    "45": "45-49",
    "50": "50-54",
    "55": "55-59",
    "60": "60-64",
    "65": "65-69",
    "70": "70-74",
    "75": "75-79",
    "80": "80-84",
    "85": "85-89",
    "90": "90+",
}


def clean_code(values: pd.Series) -> pd.Series:
    return values.astype("string").str.replace("'", "", regex=False).str.strip()


def codes_for_year(kosis_codes: pd.Series, year: int) -> pd.Series:
    codes = kosis_codes.astype("string").copy()
    if year <= 2017:
        codes = codes.replace({"23090": "23030"})
    return codes


def read_zip_csv(path: Path, encoding: str) -> pd.DataFrame:
    with zipfile.ZipFile(path) as archive:
        members = [name for name in archive.namelist() if name.lower().endswith(".csv")]
        if len(members) != 1:
            raise RuntimeError(f"Expected one CSV in {path.name}, found {members}")
        payload = archive.read(members[0])
    return pd.read_csv(io.BytesIO(payload), skiprows=2, encoding=encoding, dtype=str)


def load_boundaries() -> gpd.GeoDataFrame:
    archive = (RAW_DIR / "SIG_201905.zip").resolve()
    source = "zip://" + str(archive).replace("\\", "/")
    boundaries = gpd.read_file(source, encoding="cp949")
    boundaries = boundaries[["SIG_CD", "SIG_ENG_NM", "SIG_KOR_NM", "geometry"]].copy()
    boundaries = boundaries.sort_values("SIG_CD").reset_index(drop=True)
    return boundaries


def build_contiguity(boundaries: gpd.GeoDataFrame) -> tuple[np.ndarray, list[set[int]]]:
    graph = nx.Graph()
    graph.add_nodes_from(range(len(boundaries)))
    spatial_index = boundaries.sindex

    for left, geometry in enumerate(boundaries.geometry):
        for right in spatial_index.query(geometry, predicate="intersects"):
            right = int(right)
            if right > left and geometry.boundary.intersects(
                boundaries.geometry.iloc[right].boundary
            ):
                graph.add_edge(left, right)

    adjacency = nx.to_numpy_array(graph, nodelist=range(len(boundaries)), dtype=np.int8)
    components = sorted(nx.connected_components(graph), key=len, reverse=True)
    return adjacency, components


def build_kosis_crosswalk(boundaries: gpd.GeoDataFrame) -> pd.DataFrame:
    population = read_zip_csv(RAW_DIR / "kosis_midyear_population_2019.zip", "cp949")
    regions = population.iloc[:, :2].drop_duplicates().copy()
    regions.columns = ["kosis_code", "kosis_name_korean"]
    regions["kosis_code"] = clean_code(regions["kosis_code"])
    regions["kosis_name_korean"] = regions["kosis_name_korean"].str.strip()
    regions = regions[regions["kosis_code"].str.len() == 5].copy()
    regions["province_code"] = regions["kosis_code"].str[:2].map(OLD_TO_CURRENT_PROVINCE)
    regions["name_key"] = regions["kosis_name_korean"].str.replace(" ", "", regex=False)

    crosswalk_rows = []
    for row in boundaries.itertuples():
        province_code = row.SIG_CD[:2]
        name_key = row.SIG_KOR_NM.replace(" ", "")
        candidates = regions[
            (regions["province_code"] == province_code) & (regions["name_key"] == name_key)
        ]
        if candidates.empty:
            candidates = regions[
                (regions["province_code"] == province_code)
                & regions["name_key"].map(name_key.endswith)
            ]
        if candidates.empty and row.SIG_CD == "36110":
            candidates = regions[regions["kosis_code"] == "29010"]
        if len(candidates) != 1:
            raise RuntimeError(
                f"Could not uniquely map {row.SIG_CD} {row.SIG_KOR_NM}: "
                f"{candidates[['kosis_code', 'kosis_name_korean']].to_dict('records')}"
            )
        match = candidates.iloc[0]
        crosswalk_rows.append(
            {
                "area_id": row.SIG_CD,
                "kosis_code": match["kosis_code"],
                "kosis_name_korean": match["kosis_name_korean"],
            }
        )

    crosswalk = pd.DataFrame(crosswalk_rows)
    if not crosswalk["kosis_code"].is_unique:
        raise RuntimeError("Boundary-to-KOSIS crosswalk is not one-to-one.")
    return crosswalk


def load_observed_counts(kosis_codes: pd.Series) -> dict[str, np.ndarray]:
    totals = {key: np.zeros(len(kosis_codes), dtype=float) for key in CAUSES}
    for year in YEARS:
        year_codes = codes_for_year(kosis_codes, year)
        frame = read_zip_csv(RAW_DIR / f"kosis_municipal_mortality_{year}.zip", "utf-8-sig")
        cause_name, region_code, sex_name, deaths = frame.columns[[1, 2, 5, 7]]
        frame[cause_name] = frame[cause_name].str.strip()
        frame[region_code] = clean_code(frame[region_code])
        frame[sex_name] = frame[sex_name].str.strip()
        frame[deaths] = pd.to_numeric(frame[deaths], errors="coerce")

        for key, label in CAUSES.items():
            selected = frame[(frame[cause_name] == label) & (frame[sex_name] == "Total")]
            counts = selected.set_index(region_code)[deaths].reindex(year_codes)
            if counts.isna().any():
                missing = counts.index[counts.isna()].tolist()
                raise RuntimeError(f"Missing {key} counts in {year}: {missing}")
            totals[key] += counts.to_numpy(float)
    return totals


def prepare_population(frame: pd.DataFrame) -> pd.DataFrame:
    region_code, sex_code, age_code, population = frame.columns[[0, 2, 4, 7]]
    prepared = frame[[region_code, sex_code, age_code, population]].copy()
    prepared.columns = ["region_code", "sex", "age_code", "population"]
    prepared["region_code"] = clean_code(prepared["region_code"])
    prepared["sex"] = clean_code(prepared["sex"])
    prepared["age_code"] = clean_code(prepared["age_code"])
    prepared["age"] = prepared["age_code"].map(POPULATION_AGE_MAP)
    prepared["population"] = pd.to_numeric(prepared["population"], errors="coerce")
    prepared = prepared[
        prepared["sex"].isin(["1", "2"])
        & prepared["age"].notna()
        & prepared["population"].notna()
    ]
    return (
        prepared.groupby(["region_code", "sex", "age"], as_index=False)["population"]
        .sum()
    )


def prepare_national_deaths(frame: pd.DataFrame, cause_label: str) -> pd.DataFrame:
    cause_name, age_code, sex_code, deaths = frame.columns[[1, 4, 6, 9]]
    prepared = frame[[cause_name, age_code, sex_code, deaths]].copy()
    prepared.columns = ["cause", "age_code", "sex", "deaths"]
    prepared["cause"] = prepared["cause"].str.strip()
    prepared["age_code"] = clean_code(prepared["age_code"])
    prepared["sex"] = clean_code(prepared["sex"])
    prepared["age"] = prepared["age_code"].map(MORTALITY_AGE_MAP)
    prepared["deaths"] = pd.to_numeric(prepared["deaths"], errors="coerce")
    prepared = prepared[
        (prepared["cause"] == cause_label)
        & prepared["sex"].isin(["1", "2"])
        & prepared["age"].notna()
        & prepared["deaths"].notna()
    ]
    return prepared.groupby(["sex", "age"], as_index=False)["deaths"].sum()


def load_expected_counts(kosis_codes: pd.Series) -> dict[str, np.ndarray]:
    totals = {key: np.zeros(len(kosis_codes), dtype=float) for key in CAUSES}
    for year in YEARS:
        year_codes = codes_for_year(kosis_codes, year)
        population_raw = read_zip_csv(
            RAW_DIR / f"kosis_midyear_population_{year}.zip", "cp949"
        )
        national_mortality = read_zip_csv(
            RAW_DIR / f"kosis_national_mortality_{year}.zip", "utf-8-sig"
        )
        population = prepare_population(population_raw)
        national_population = population[population["region_code"] == "00"]
        municipal_population = population[population["region_code"].isin(year_codes)]

        expected_keys = pd.MultiIndex.from_product(
            [year_codes, ["1", "2"], sorted(set(POPULATION_AGE_MAP.values()))],
            names=["region_code", "sex", "age"],
        )
        municipal_population = (
            municipal_population.set_index(["region_code", "sex", "age"])
            .reindex(expected_keys)
            .reset_index()
        )
        if municipal_population["population"].isna().any():
            missing = municipal_population.loc[
                municipal_population["population"].isna(), ["region_code", "sex", "age"]
            ]
            raise RuntimeError(f"Missing population cells in {year}: {missing.to_dict('records')}")

        for key, label in CAUSES.items():
            deaths = prepare_national_deaths(national_mortality, label)
            rates = deaths.merge(national_population, on=["sex", "age"], how="left")
            rates["rate"] = rates["deaths"] / rates["population"]
            expected = municipal_population.merge(
                rates[["sex", "age", "rate"]], on=["sex", "age"], how="left"
            )
            if expected["rate"].isna().any():
                raise RuntimeError(f"Missing national {key} mortality rates in {year}.")
            annual = (
                expected.assign(expected=expected["population"] * expected["rate"])
                .groupby("region_code")["expected"]
                .sum()
                .reindex(year_codes)
            )
            totals[key] += annual.to_numpy(float)
    return totals


def load_kdca_indicator(
    kosis_codes: pd.Series, indicator: str
) -> tuple[np.ndarray, list[str]]:
    workbook = openpyxl.load_workbook(
        RAW_DIR / "kdca_health_database_v1_7.xlsx", read_only=True, data_only=True
    )
    annual = []
    sources_by_year: list[list[str]] = []
    deprivation_term = "\uae30\ucd08\uc0dd\ud65c"
    smoking_name = "\ud604\uc7ac\ud761\uc5f0\uc728_\ud45c\uc900\ud654\uc728"

    for year in YEARS:
        # KDCA retains the pre-rename Michuhol-gu code throughout this workbook.
        year_codes = kosis_codes.astype("string").replace({"23090": "23030"})
        worksheet = workbook[str(year)]
        header = list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        if indicator == "deprivation":
            columns = [
                index
                for index, value in enumerate(header)
                if value and deprivation_term in str(value)
            ]
        elif indicator == "smoking":
            columns = [
                index for index, value in enumerate(header) if str(value) == smoking_name
            ]
        else:
            raise ValueError(f"Unknown KDCA indicator: {indicator}")
        if not columns:
            raise RuntimeError(f"Could not locate {indicator} in the {year} worksheet.")
        value_column = columns[-1]
        values: dict[str, float] = {}
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            code = row[1]
            value = row[value_column]
            if code is None or value in (None, ""):
                continue
            values[str(int(code))] = float(value)

        year_values = []
        year_sources = []
        for code in year_codes:
            candidates = [code, code[:4] + "0"]
            if code == "29010":
                candidates.append("29")
            source = next((candidate for candidate in candidates if candidate in values), None)
            if source is None:
                raise RuntimeError(
                    f"Missing {indicator} value for KOSIS code {code} in {year}."
                )
            year_values.append(values[source])
            year_sources.append(source)
        annual.append(year_values)
        sources_by_year.append(year_sources)

    workbook.close()
    source_codes = [
        "/".join(dict.fromkeys(year_sources[index] for year_sources in sources_by_year))
        for index in range(len(kosis_codes))
    ]
    return np.mean(np.asarray(annual, dtype=float), axis=0), source_codes


def write_outputs() -> None:
    boundaries = load_boundaries()
    adjacency_all, components = build_contiguity(boundaries)
    if len(components[0]) != 241:
        raise RuntimeError(
            f"Expected a 241-area mainland component, found {len(components[0])}."
        )
    boundaries["geometry"] = boundaries.geometry.map(make_valid)
    largest_component = sorted(components[0])
    included = np.zeros(len(boundaries), dtype=bool)
    included[largest_component] = True

    crosswalk = build_kosis_crosswalk(boundaries)
    data = boundaries.merge(crosswalk, left_on="SIG_CD", right_on="area_id", validate="one_to_one")
    kosis_codes = data["kosis_code"]
    observed = load_observed_counts(kosis_codes)
    expected = load_expected_counts(kosis_codes)
    deprivation, deprivation_sources = load_kdca_indicator(kosis_codes, "deprivation")
    smoking, smoking_sources = load_kdca_indicator(kosis_codes, "smoking")

    data = data.rename(
        columns={
            "SIG_CD": "boundary_code",
            "SIG_ENG_NM": "area_name",
            "SIG_KOR_NM": "area_name_korean",
        }
    )
    data["analysis_included"] = included
    data["component_id"] = -1
    for component_id, component in enumerate(components):
        data.loc[list(component), "component_id"] = component_id
    data["degree"] = adjacency_all.sum(axis=1)
    data["deprivation_pct"] = deprivation
    data["deprivation_source_code"] = deprivation_sources
    data["smoking_pct"] = smoking
    data["smoking_source_code"] = smoking_sources

    for key in CAUSES:
        data[f"observed_{key}"] = observed[key]
        data[f"expected_{key}"] = expected[key]
        data[f"smr_{key}"] = observed[key] / expected[key]

    analysis = data.loc[included].copy().reset_index(drop=True)
    deprivation_mean = analysis["deprivation_pct"].mean()
    deprivation_sd = analysis["deprivation_pct"].std(ddof=0)
    smoking_mean = analysis["smoking_pct"].mean()
    smoking_sd = analysis["smoking_pct"].std(ddof=0)
    data["deprivation_std"] = (data["deprivation_pct"] - deprivation_mean) / deprivation_sd
    data["smoking_std"] = (data["smoking_pct"] - smoking_mean) / smoking_sd
    analysis = data.loc[included].copy().reset_index(drop=True)

    # Generic aliases use the primary lung-cancer/current-smoking application.
    analysis["observed"] = analysis["observed_lung_cancer"]
    analysis["expected"] = analysis["expected_lung_cancer"]
    analysis["SMR"] = analysis["smr_lung_cancer"]
    analysis["risk"] = analysis["smr_lung_cancer"]

    adjacency = adjacency_all[np.ix_(largest_component, largest_component)]
    adjacency_frame = pd.DataFrame(
        adjacency,
        index=analysis["area_id"].astype(str),
        columns=analysis["area_id"].astype(str),
    )
    adjacency_frame.to_csv(BASE_DIR / "adjacency_matrix_south_korea.csv", index=False)

    data.to_file(BASE_DIR / "mortality_data_south_korea_all.gpkg", driver="GPKG")
    analysis.to_file(BASE_DIR / "mortality_data_south_korea.gpkg", driver="GPKG")
    crosswalk.to_csv(BASE_DIR / "south_korea_kosis_crosswalk.csv", index=False)

    excluded = data.loc[~included, ["area_id", "area_name", "area_name_korean"]]
    excluded.to_csv(BASE_DIR / "excluded_island_municipalities.csv", index=False)

    edge_rows, edge_columns = np.where(np.triu(adjacency, k=1) == 1)
    covariate_metrics = {}
    for key in ["smoking", "deprivation"]:
        x = analysis[f"{key}_std"].to_numpy(float)
        edge_dissimilarity = np.abs(x[edge_rows] - x[edge_columns])
        x_centered = x - x.mean()
        moran_i = (
            len(x)
            / adjacency.sum()
            * (x_centered @ adjacency @ x_centered)
            / (x_centered @ x_centered)
        )
        covariate_metrics[key] = {
            "mean_pct": float(analysis[f"{key}_pct"].mean()),
            "sd_pct": float(analysis[f"{key}_pct"].std(ddof=0)),
            "median_standardized_edge_dissimilarity": float(
                np.median(edge_dissimilarity)
            ),
            "M": float(np.log(2.0) / np.median(edge_dissimilarity)),
            "moran_i": float(moran_i),
        }
    manifest = {
        "period": "2015-2019",
        "full_area_count": int(len(data)),
        "analysis_area_count": int(len(analysis)),
        "component_sizes": [int(len(component)) for component in components],
        "analysis_edge_count": int(adjacency.sum() // 2),
        "analysis_mean_degree": float(adjacency.sum(axis=1).mean()),
        "primary_application": "lung_cancer_with_current_smoking",
        "covariates": covariate_metrics,
        "outcomes": {
            key: {
                "cause": label,
                "observed_total": int(analysis[f"observed_{key}"].sum()),
                "expected_total": float(analysis[f"expected_{key}"].sum()),
                "observed_range": [
                    int(analysis[f"observed_{key}"].min()),
                    int(analysis[f"observed_{key}"].max()),
                ],
                "expected_range": [
                    float(analysis[f"expected_{key}"].min()),
                    float(analysis[f"expected_{key}"].max()),
                ],
            }
            for key, label in CAUSES.items()
        },
    }
    (BASE_DIR / "south_korea_data_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )

    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    write_outputs()
